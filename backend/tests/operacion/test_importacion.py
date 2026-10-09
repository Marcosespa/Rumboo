import re
import time
import zipfile
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from app.operacion import excel
from app.operacion.models import Conductor
from tests.ayudas import trip_body

XLSX = {"Content-Type": excel.XLSX}


def fila(manifiesto, placa="ABC123", cedula=1012345678, remesa="R-1", peso=1000, **cambios):
    return {"manifiesto": manifiesto, "origen": "Bogotá", "destino": "Medellín",
            "salida_estimada": datetime(2026, 12, 1, 8, 0), "llegada_estimada": datetime(2026, 12, 1, 18, 0),
            "conductor_nombre": "Ana Ruiz", "conductor_cedula": cedula, "conductor_telefono": "3001234567",
            "placa": placa, "propietario": "Flota", "remesa_numero": remesa, "remesa_cliente": "Cliente",
            "remesa_peso_kg": peso, "remesa_cantidad": 10} | cambios


def libro(filas, columnas=tuple(excel.COLUMNAS)):
    wb = Workbook()
    hoja = wb.active
    hoja.append(list(columnas))
    for f in filas:
        hoja.append([f.get(c) for c in columnas])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def importar(client, headers, contenido):
    return client.post("/api/viajes/importar", content=contenido, headers=headers | XLSX)


def test_import_creates_one_trip_per_manifest_with_colombian_dates(client, login):
    headers = login()
    archivo = libro([fila("MAN-A"), fila("MAN-B", placa="XYZ789", cedula=1099999999), fila("MAN-A", remesa="R-2", peso=500)])
    body = importar(client, headers, archivo).json()
    assert [c["manifiesto"] for c in body["creados"]] == ["MAN-A", "MAN-B"] and body["errores"] == []
    trip = client.get(f"/api/viajes/{body['creados'][0]['id']}", headers=headers).json()
    assert trip["peso_salida_kg"] == 1500 and [r["numero"] for r in trip["remesas"]] == ["R-1", "R-2"]
    assert trip["salida_estimada"].startswith("2026-12-01T13:00:00")  # 08:00 en Bogotá
    assert trip["estado"] == "registrado" and trip["conductor"]["autoriza_contacto"] is False


def test_invalid_group_creates_nothing_and_does_not_block_the_others(client, login):
    headers = login()
    archivo = libro([fila("MAN-A", placa="PLACA-SECRETA"), fila("MAN-A", remesa="R-2", placa="PLACA-SECRETA"),
                     fila("MAN-B", placa="XYZ789", cedula=1099999999),
                     fila("MAN-C", placa="DEF456", cedula=1088888888, peso=-5),
                     fila("MAN-D", placa="GHI012", cedula=1077777777, destino="Cali"),
                     fila("MAN-D", placa="GHI012", cedula=1077777777, remesa="R-2", destino="Cartagena")])
    response = importar(client, headers, archivo)
    body = response.json()
    assert [c["manifiesto"] for c in body["creados"]] == ["MAN-B"]
    errores = {e["manifiesto"]: e for e in body["errores"]}
    assert errores["MAN-A"]["campo"] == "placa" and errores["MAN-A"]["filas"] == [2, 3]
    assert errores["MAN-C"]["campo"] == "remesa_peso_kg" and errores["MAN-C"]["filas"] == [5]
    assert errores["MAN-D"]["campo"] == "destino"
    assert "PLACA-SECRETA" not in response.text
    assert client.get("/api/viajes", headers=headers).json()["total"] == 1


def test_reimporting_the_same_file_creates_no_duplicates(client, login):
    headers = login()
    archivo = libro([fila("MAN-A")])
    assert len(importar(client, headers, archivo).json()["creados"]) == 1
    again = importar(client, headers, archivo).json()
    assert again["creados"] == []
    assert again["errores"][0]["mensaje"] == "Ya existe un viaje con ese manifiesto"
    assert client.get("/api/viajes", headers=headers).json()["total"] == 1


def test_import_reuses_existing_catalog_without_changing_it(app, client, login):
    headers = login()
    trip = client.post("/api/viajes", json=trip_body(), headers=headers).json()
    client.post(f"/api/viajes/{trip['id']}/transiciones", json={"estado": "cancelado", "motivo": "Prueba"}, headers=headers)
    with app.state.sessions() as db:
        authorized_at = db.scalar(select(Conductor.autorizado_en))
    other_phone = importar(client, headers, libro([fila("MAN-Y", conductor_telefono="3109876543")])).json()
    assert other_phone["creados"] == [] and "celular" in other_phone["errores"][0]["mensaje"]
    same = libro([fila("MAN-X", conductor_nombre="Otro Nombre", propietario=None)])
    assert importar(client, headers, same).json()["errores"] == []
    with app.state.sessions() as db:
        driver = db.scalar(select(Conductor))
        assert driver.autoriza_contacto and driver.autorizado_en == authorized_at
        assert driver.nombre == "Juan Pérez" and driver.telefono == "+573001234567"
    assert client.get("/api/vehiculos", headers=headers).json()["items"][0]["propietario"] == "Flota"


def test_numbers_are_not_accepted_as_dates(client, login):
    headers = login()
    body = importar(client, headers, libro([fila("MAN-N", salida_estimada=46357.33, llegada_estimada="46358")])).json()
    assert body["creados"] == []
    assert {e["campo"] for e in body["errores"]} == {"salida_estimada", "llegada_estimada"}


def test_import_is_isolated_per_carrier(client, login):
    first, second = login(), login("Otra Empresa", "otro")
    archivo = libro([fila("MAN-A")])
    created = importar(client, first, archivo).json()["creados"][0]
    assert importar(client, second, archivo).json()["creados"][0]["manifiesto"] == "MAN-A"
    assert client.get("/api/viajes", headers=first).json()["total"] == 1
    assert client.get(f"/api/viajes/{created['id']}", headers=second).status_code == 404


def test_invalid_or_oversized_files_are_rejected(client, login):
    headers = login()
    assert importar(client, headers, b"no es un excel").json()["detail"] == "El archivo debe ser un Excel .xlsx"
    missing = importar(client, headers, libro([], columnas=("manifiesto", "origen")))
    assert missing.status_code == 422 and "placa" in missing.json()["detail"]
    assert importar(client, headers, b"x" * (excel.MAX_BYTES + 1)).status_code == 413
    assert client.post("/api/viajes/importar", content=libro([fila("MAN-A")]), headers=XLSX).status_code == 401


def reemplazar(contenido, nombre, cambio):
    """Copia el .xlsx aplicando `cambio` a una de sus partes."""
    origen, destino = zipfile.ZipFile(BytesIO(contenido)), BytesIO()
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zf:
        for info in origen.infolist():
            data = origen.read(info.filename)
            zf.writestr(info, cambio(data) if info.filename == nombre else data)
    return destino.getvalue()


def test_oversized_part_is_rejected_before_parsing(client, login, monkeypatch):
    headers = login()
    monkeypatch.setattr(excel, "load_workbook", lambda *a, **k: (_ for _ in ()).throw(AssertionError("no debe abrirse")))
    # Una sola fila con millones de celdas vacías: pocos KB comprimidos, cientos de MB al parsear.
    wide = reemplazar(libro([]), "xl/worksheets/sheet1.xml", lambda _: b"<c/>" * (excel.MAX_HOJA // 4 + 1))
    assert len(wide) < excel.MAX_BYTES
    response = importar(client, headers, wide)
    assert response.status_code == 422 and response.json()["detail"] == "El archivo es demasiado grande para importarlo"


def test_tampered_file_gives_422_without_logging_cell_values(client, login, caplog):
    headers = login()
    tampered = reemplazar(libro([fila("MAN-A", peso=4321)]), "xl/worksheets/sheet1.xml",
                          lambda xml: xml.replace(b"<v>4321</v>", b"<v>CeldaPrivada</v>"))
    response = importar(client, headers, tampered)
    assert response.status_code == 422 and response.json()["detail"] == "El archivo no se pudo leer; usa la plantilla de Rumboo"
    assert "CeldaPrivada" not in caplog.text and "CeldaPrivada" not in response.text


def test_declared_huge_dimension_does_not_widen_rows(client, login):
    headers = login()
    huge = reemplazar(libro([]), "xl/worksheets/sheet1.xml", lambda xml: re.sub(
        rb'<dimension ref="[^"]*" ?/>', b'<dimension ref="A1:ZZZ2001"/>', xml).replace(b"</sheetData>", b'<row r="2001"/></sheetData>'))
    assert b"ZZZ2001" in zipfile.ZipFile(BytesIO(huge)).read("xl/worksheets/sheet1.xml")
    started = time.monotonic()
    assert excel.leer_filas(huge) == []
    # Sin acotar columnas son 2001 filas de 18 278 celdas (~0,9 s de CPU); acotadas, milisegundos.
    assert time.monotonic() - started < 0.3
    assert importar(client, headers, huge).json() == {"creados": [], "errores": []}


def test_concurrent_import_is_refused_while_another_runs(client, login):
    headers = login()
    assert excel._lectura.acquire(blocking=False)
    try:
        response = importar(client, headers, libro([fila("MAN-A")]))
    finally:
        excel._lectura.release()
    assert response.status_code == 409 and "otra importación" in response.json()["detail"]


def test_template_has_the_import_columns(client, login):
    headers = login()
    response = client.get("/api/viajes/plantilla", headers=headers)
    assert response.status_code == 200 and response.headers["content-type"] == excel.XLSX
    header = next(load_workbook(BytesIO(response.content), read_only=True).worksheets[0].iter_rows(values_only=True))
    assert list(header) == list(excel.COLUMNAS)
    assert importar(client, headers, response.content).json() == {"creados": [], "errores": []}
    assert client.get("/api/viajes/plantilla").status_code == 401
