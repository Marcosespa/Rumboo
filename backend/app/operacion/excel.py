"""Lectura y plantilla del Excel de viajes. Solo traduce el archivo a filas; las reglas viven en servicio.py.

El .xlsx es un zip de XML: antes de abrirlo se acotan entradas y tamaño descomprimido por parte, porque openpyxl
arma cada fila completa en memoria (una hoja de 4 MB puede pedir ~300 MB). Se lee en modo solo lectura, con valores
guardados (sin evaluar fórmulas ni macros), columnas acotadas y una importación a la vez por proceso.
"""
import threading
import zipfile
import zlib
from io import BytesIO
from xml.etree.ElementTree import ParseError
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from app.core.errores import Conflicto, Invalido

MAX_BYTES = 1_000_000
MAX_HOJA = 4_000_000  # una hoja legítima de 2000 filas ocupa ~1,6 MB
MAX_PARTE = 1_000_000
MAX_DESCOMPRIMIDO = 10_000_000
MAX_ENTRADAS = 200
MAX_FILAS = 2000
_lectura = threading.BoundedSemaphore(1)
COLUMNAS = {
    "manifiesto": "Número de manifiesto; las filas con el mismo número forman un viaje",
    "origen": "Ciudad de origen",
    "destino": "Ciudad de destino",
    "salida_estimada": "Fecha y hora de salida; sin zona se toma la hora de Colombia",
    "llegada_estimada": "Fecha y hora de llegada; posterior a la salida",
    "conductor_nombre": "Nombre del conductor",
    "conductor_cedula": "Cédula del conductor, 6 a 10 dígitos",
    "conductor_telefono": "Celular colombiano de 10 dígitos",
    "placa": "Placa del vehículo, tres letras y tres números",
    "propietario": "Propietario del vehículo (opcional)",
    "remesa_numero": "Número de remesa, único dentro del manifiesto",
    "remesa_cliente": "Cliente de la remesa",
    "remesa_peso_kg": "Peso de la remesa en kilogramos",
    "remesa_cantidad": "Unidades de la remesa (opcional)",
}
OPCIONALES = {"propietario", "remesa_cantidad"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _validar_zip(contenido):
    try:
        entradas = zipfile.ZipFile(BytesIO(contenido)).infolist()
    except zipfile.BadZipFile as err:
        raise Invalido("El archivo debe ser un Excel .xlsx") from err
    grandes = ("xl/worksheets/", "xl/sharedStrings.xml")
    if (len(entradas) > MAX_ENTRADAS or sum(e.file_size for e in entradas) > MAX_DESCOMPRIMIDO
            or any(e.file_size > (MAX_HOJA if e.filename.startswith(grandes) else MAX_PARTE) for e in entradas)):
        raise Invalido("El archivo es demasiado grande para importarlo")


# Errores de openpyxl ante un .xlsx manipulado; su mensaje puede repetir celdas, por eso no se propaga.
_ARCHIVO_DANADO = (InvalidFileException, KeyError, ParseError, zipfile.BadZipFile, zlib.error, ValueError, TypeError,
                   IndexError, AttributeError, OverflowError)


def leer_filas(contenido):
    """Devuelve dicts con las columnas conocidas y `fila` (número en Excel). Omite filas vacías."""
    _validar_zip(contenido)
    if not _lectura.acquire(blocking=False):
        raise Conflicto("Hay otra importación en curso; intenta de nuevo en unos segundos")
    try:
        return _leer(contenido)
    except _ARCHIVO_DANADO as err:
        raise Invalido("El archivo no se pudo leer; usa la plantilla de Rumboo") from err
    finally:
        _lectura.release()


def _leer(contenido):
    libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    try:
        # max_col evita que una dimensión declarada enorme multiplique el trabajo por fila.
        filas = libro.worksheets[0].iter_rows(max_col=len(COLUMNAS) + 16, values_only=True)
        encabezado = [str(c).strip().lower() if c is not None else "" for c in next(filas, ())]
        faltantes = [c for c in COLUMNAS if c not in encabezado and c not in OPCIONALES]
        if faltantes:
            raise Invalido(f"Faltan columnas en la plantilla: {', '.join(faltantes)}")
        resultado = []
        for numero, valores in enumerate(filas, start=2):
            if numero - 1 > MAX_FILAS:
                raise Invalido(f"El archivo supera {MAX_FILAS} filas")
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in valores):
                continue
            fila = {col: val for col, val in zip(encabezado, valores, strict=False) if col in COLUMNAS}
            resultado.append(fila | {"fila": numero})
        return resultado
    finally:
        libro.close()


def plantilla():
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Viajes"
    hoja.append(list(COLUMNAS))
    instrucciones = libro.create_sheet("Instrucciones")
    instrucciones.append(["Columna", "Contenido"])
    for columna, descripcion in COLUMNAS.items():
        instrucciones.append([columna, descripcion])
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
